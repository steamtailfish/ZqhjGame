"""Build the technical report from the frozen v22 evidence."""
from pathlib import Path
import json,hashlib
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image,ImageDraw,ImageFont

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'artifacts/reports/v22-technical'
OUT.mkdir(parents=True,exist_ok=True)
PACKAGE=ROOT/'artifacts/submission/score-v22'
M=json.loads((PACKAGE/'manifest.json').read_text(encoding='utf-8'))
E=json.loads((PACKAGE/'baseline_evaluation.json').read_text(encoding='utf-8'))
assert E['total_score']==9.06 and E['n_reports']==2
for name,key in [('agent.py','agent_sha256'),('vision.pt','vision_sha256'),('baseline_evaluation.json','evaluation_sha256')]:
    assert hashlib.sha256((PACKAGE/name).read_bytes()).hexdigest()==M[key]

font=lambda n:ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',n)
im=Image.new('RGB',(1400,420),'white');d=ImageDraw.Draw(im)
def box(rect,lines):
    d.rounded_rectangle(rect,12,fill='#F0F4F7',outline='#687A87',width=2)
    x1,y1,x2,y2=rect
    for i,line in enumerate(lines):
        b=d.textbbox((0,0),line,font=font(25));d.text(((x1+x2-b[2])/2,y1+20+i*36),line,font=font(25),fill='black')
def arrow(points):
    d.line(points,fill='#405363',width=3)
    x,y=points[-1];px,py=points[-2]
    if x>px:d.polygon([(x,y),(x-12,y-7),(x-12,y+7)],fill='#405363')
    elif x<px:d.polygon([(x,y),(x+12,y-7),(x+12,y+7)],fill='#405363')
    else:d.polygon([(x,y),(x-7,y-12),(x+7,y-12)],fill='#405363')
box((20,30,300,145),['本机公开照片','1024 × 768 RGB'])
box((365,30,640,145),['轮廓候选与分类','64 × 64 私有网络'])
box((705,30,980,145),['跨帧关联','相对背景运动'])
box((1045,30,1380,145),['本机位姿与局部平面','估计地理坐标及预算'])
for a,b in [(300,365),(640,705),(980,1045)]:arrow([(a,87),(b,87)])
box((1010,260,1380,375),['身份 时效 运动门限','坐标上报 最高 1 Hz'])
box((355,260,850,375),['合法广播 协同分工','势场引导与解析运动基元规划'])
box((20,260,300,375),['飞行与云台动作','下一时刻重新观测'])
arrow([(1210,145),(1210,260)])
arrow([(1120,145),(1120,205),(605,205),(605,260)])
arrow([(355,315),(300,315)])
im.save(OUT/'architecture.png')

im=Image.new('RGB',(1400,440),'white');d=ImageDraw.Draw(im)
x0,x1,y0,y1=100,1280,40,350
for s in (0,2,4,6,8,10):
    y=y1-(y1-y0)*s/10;d.line((x0,y,x1,y),fill='#DDDDDD');d.text((40,y-16),str(s),font=font(23),fill='black')
for t in (0,100,200,300,400,500,600):
    x=x0+(x1-x0)*t/600;d.text((x-20,y1+12),str(t),font=font(23),fill='black')
points=[(x0+(x1-x0)*r['sim_time']/600,y1-(y1-y0)*r['total_score']/10) for r in E['score_timeline']]
d.line(points,fill='#254C68',width=4)
d.text((1150,18),'最终 9.06',font=font(26),fill='black')
d.text((520,392),'仿真时间 秒',font=font(25),fill='black');d.text((10,2),'总分',font=font(23),fill='black')
im.save(OUT/'official-score.png')

doc=Document();sec=doc.sections[0]
for style in doc.styles:
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2);sec.bottom_margin=Cm(1.9);sec.left_margin=Cm(2.15);sec.right_margin=Cm(2.15)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    s=doc.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
    s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体' if name=='Normal' else '微软雅黑')
    s.paragraph_format.space_after=Pt(7)
doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.line_spacing=1.22
doc.styles['Title'].font.size=Pt(22);doc.styles['Title'].font.bold=True
doc.styles['Subtitle'].font.size=Pt(11)
doc.styles['Subtitle'].font.italic=False
for name,size in [('Heading 1',16),('Heading 2',12)]:
    doc.styles[name].font.size=Pt(size);doc.styles[name].font.bold=True
    doc.styles[name].paragraph_format.space_before=Pt(10);doc.styles[name].paragraph_format.keep_with_next=True
doc.styles['Caption'].font.size=Pt(9);doc.styles['Caption'].paragraph_format.space_after=Pt(9)
doc.core_properties.title='多无人机视觉协同定位与上报方法'
doc.core_properties.subject='ZqhjGame score v22 当前实现与正式实验结果'
doc.core_properties.author='ZqhjGame 项目组'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('第 ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
footer.add_run(' 页')
for r in footer.runs:r.font.size=Pt(9)
def p(t,style=None):return doc.add_paragraph(t,style)
def h(t,level=1):return doc.add_heading(t,level)
def page(t):doc.add_page_break();h(t)
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for i,txt in enumerate(headers):t.rows[0].cells[i].text=txt
    for row in rows:
        cells=t.add_row().cells
        for c,txt in zip(cells,row):c.text=str(txt)
    pr=t._tbl.tblPr;b=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
    pr.append(b)
    for ri,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                v=OxmlElement('w:'+side);v.set(qn('w:w'),'90');v.set(qn('w:type'),'dxa');marg.append(v)
            cp.append(marg)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E5EDF3');cp.append(sh)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.05
                if ci==0:para.alignment=WD_ALIGN_PARAGRAPH.LEFT
                for r in para.runs:r.font.size=Pt(9.5);r.font.bold=(ri==0)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
    p('').paragraph_format.space_after=Pt(0)
    return t
def equation(text):
    para=p('');para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    om=OxmlElement('m:oMath');r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);om.append(r);para._p.append(om)
def picture(name,width=16.5):
    para=p('');para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=para.add_run();r.add_picture(str(OUT/name),width=Cm(width))
    dp=r._r.xpath('.//wp:docPr')[0];dp.set('descr','系统处理流程' if name=='architecture.png' else '官方总分随仿真时间的变化')
def code(text):
    para=p('');para.paragraph_format.space_after=Pt(8);para.paragraph_format.line_spacing=1.0
    r=para.add_run(text);r.font.name='Consolas';r.font.size=Pt(8)
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')

p('多无人机视觉协同定位与上报方法','Title')
p('ZqhjGame score v22 技术报告','Subtitle')
p('版本日期 2026年9月11日    适用任务 红枫2026赛题二')
h('1 当前方法与阶段结论')
p('我们当前采用模块化方法：先从本机RGB照片中提取车辆候选，用轻量卷积网络判断外观类别，再结合连续图像运动和本机位姿估计目标坐标。飞行侧通过公开区域搜索、合法广播和解析轨迹规划执行协同动作；上报侧独立检查身份、运动、定位一致性与时效，满足条件才报告坐标。')
p('score v22已在原始正式场景取得9.06分。两次坐标报告的累计定位RMSE为11.27米，没有惩罚；三个目标均未清除，实际同目标双机有效跟踪计数均为0。因此，本阶段完成了视觉定位到正式精度得分的闭环，尚未完成持续协同捕获。正式结果与冻结代码见文末证据索引。[3][4]')
table(['正式总分','坐标报告','目标清除','惩罚'],[['9.06','2次','0 / 3','0']], [4.1,4.1,4.1,4.1])
picture('architecture.png')
p('图1 当前处理流程  坐标报告不等待导航轨迹确认','Caption')
h('方法定位',2)
p('本方案参考Gou等的策略引导围捕思想，以及YOPO的运动基元与轨迹代价思想。[1][2] 当前启用的学习模块是外观分类网络；规划端实际运行解析代价计算。此前实现的YOPO式引导学习代码仍保留，但学习评分头未用于这次拿分版本，图像到轨迹之间也没有统一的端到端训练计算图。')

page('2 视觉感知与训练方法')
h('轮廓候选生成',2)
p('输入为本机公开的1024×768 RGB照片。先转灰度，联合自适应阈值与35、45、55、65、75、85六个固定亮度阈值生成前景轮廓。对轮廓面积、最小外接矩形、长宽比、填充率与凸包紧致度进行筛选，并去除重叠候选。面积范围为10至1600像素，长边8至90像素，短边至少2.5像素，长宽比1.5至4.8；候选数量上限512。')
p('每个候选按长边的1.5倍截取正方形局部区域，缩放到64×64，保留适量背景上下文并维持目标比例。网络输出真车、诱饵、背景三类概率；背景类不形成车辆输出。分类后再次以IoU 0.35去重，最多保留32个框。搜索候选门限为0.55；正式报告使用更严格的0.95置信度及0.9类别分差。')
h('局部外观网络',2)
table(['模块','结构','输出'],[
('输入','RGB局部裁剪及归一化','3 × 64 × 64'),
('卷积特征','4层3×3卷积 步长均为2\n通道数16 32 48 64 每层ReLU','64 × 4 × 4'),
('分类头','Flatten → Linear 1024到64 → ReLU\nLinear 64到3 → Softmax','真车 诱饵 背景概率')],[2.8,9.2,4.4])
h('训练配置与数据边界',2)
p('实际部署权重为appearance v3。训练使用35个真车、17个诱饵及709个背景局部样本；每批48个样本按三类均衡抽取。优化器为AdamW，学习率0.0005，权重衰减0.005，以交叉熵训练2400步。随机旋转覆盖360度，缩放范围0.8至1.15，并加入翻转、14至64像素低分辨率重采样及亮度扰动。')
p('车辆身份和框来自受控采集后的人工审核；正式照片只补充人工确认的背景。验证集由预留的受控高度布景组成，共34个样本，分类34个正确。这一结果反映有限受控数据表现，不能外推为正式场景识别率。颜色增强的appearance v4在后续复查中退步，未部署；新增待审核候选也未进入当前权重。[5]')
h('异步推理',2)
p('每架无人机初始化私有模型和单个推理工作线程，约每0.5秒提交一张新照片。图像以哈希去重，输出始终关联原照片及其接收位姿，超过1秒的异步结果丢弃；异常被捕获并返回空感知结果，防止SDK默认探测结果混入。模型在初始化加载，比赛回调不读模型文件。')

page('3 跨帧关联与地理定位')
h('补偿相机运动后的目标关联',2)
p('直接比较两帧中的像素中心，会把云台转动引起的目标位移当成丢轨。当前方法先通过背景特征估计单应变换H，将上一帧目标中心投影到当前帧，再寻找最近的候选框；关联距离限制为图像宽度的15%。缺少有效H时，用上一帧像素射线和当前本机相机姿态预测位置。')
p('背景估计采用最多250个角点、金字塔Lucas–Kanade光流和前后向检查，再经RANSAC拟合单应变换，至少需要30个内点。图像关联在云台转动时仍可进行；用于估计地面高度的几何拟合则要求最近1.5秒的光轴方位和俯仰变化不超过5度，并满足至少6米的平移基线。')
h('局部地面平面估计',2)
p('相机光轴使用机体航向加相对云台pan构造。设像素横纵坐标为u、v，主点为cx、cy，图像宽度为W，横向视场角为FOV，前、右、下三个世界坐标基向量分别为a、b、c，则：')
equation('f = W / (2 tan(FOV / 2))')
equation('ξ = (u − cx) / f     η = (v − cy) / f     r = a + ξ b + η c')
p('给定无人机位置p及高度zu，假设局部地面高度为zg，将射线与水平平面相交。水平分量按局部经纬度尺度换算为目标坐标：')
equation('λ = (zg − zu) / rz     pg = p + λ r')
p('地面高度先以20米步长粗搜索，再以2米步长细化，使两帧静态特征的重投影误差中位数较小；交替特征点用于拟合和验证。验证误差中位数须不超过4像素，90分位数不超过10像素。最终定位需至少3次平面拟合、时间跨度至少1秒，最近拟合不早于4秒，地面高度离散度不超过20米。')
h('运动证据与误差预算',2)
p('将单应变换给出的静态预测位置与车辆当前中心比较，得到相对背景残差。至少两次连续残差满足2.5像素至图像宽度12%的范围，才形成运动证据。上报分支进一步把当前中心与静态预测中心投影到同一地面平面，将距离除以帧间时间，要求估计速度在3.5至18米/秒。')
p('定位误差预算综合基础40米、水平距离、拟合高度离散度、斜视比例及姿态变化项；它是工程筛选量，不是标定后的协方差。平面重用同时受时间和120米位移限制。照片目前使用接收位姿，capture_verified仍为False；相机延迟、地形起伏、树冠和车辆自身高度会产生系统误差，不能把单次11.27米实测当成全场景精度保证。')

page('4 搜索 协同与解析规划')
h('公开区域条带搜索',2)
p('三机通过广播发现完整队伍后冻结排序，在公开任务边界内生成间隔约280米的搜索条带，以队员序号交错分配条带并按往复方式遍历。边界留出约250米余量，搜索速度设为35米/秒，FOV保持50度。路径只由公开边界、本机位置和广播队伍信息生成，不读取隐藏道路或目标路线。')
h('协同状态与通信',2)
p('地理候选进入本机alpha–beta轨迹库，用短期位置预测完成关联；轨迹至少3次命中且持续0.5秒后可参与协同。通信包采用二进制定点字段和Base85编码，基础包41字节，含地面高度与身份的扩展包45字节，低于50字节限制；发送频率控制在4Hz以内。广播包含本机位置、航向、速度及候选位置、身份、观测年龄等。')
p('协同器用稳定的拥有者与轨迹编号排序选择目标，并指定发现者及另一架无人机参与观察。候选最多短期保留8秒，一次观察尝试45秒后进入30秒冷却。OBSERVE只是控制器意图，不能替代引擎同目标双机有效跟踪判定。当前简单分工尚未解决两机对同一真车持续可见的问题。')
h('势场引导与运动基元代价',2)
p('引导方向由目标吸引、队友排斥、障碍排斥和速度阻尼组合；观察模式增加切向分量，使固定翼围绕目标运动，参考环绕半径450米。该做法借鉴策略引导围捕的结构，但参数为本比赛的工程设置，没有运行Gou等完整强化学习训练流程。[1]')
p('解析规划器以4秒为预测时域，0.25秒积分，枚举恒转率和目标速度组合。转率包括−30、−15、−5、0、5、15、30度/秒及引导方向对应的附加转率；速度候选为15、22、30米/秒及当前巡航速度。总代价为：')
equation('J = 0.08 S + 3 C + G + 2 R')
p('S为转动和加减速的平滑代价，C为随预测间距恶化而指数增加的安全代价，G为临时目标偏差，R为环绕半径偏差。优先选择满足预测间距约束的低代价候选；没有可行候选时选取最小间距最大的方案，并显式记录不可行。该检查不构成实际安全性的数学保证。')
h('动作执行与云台稳定',2)
p('每约0.5秒执行规划首段，用fly_to接口发送按规划航向构造的前方300米航点。云台跟随目标世界射线，单次pan最多改变6度、tilt最多改变3度，避免重复积分同一旧照片的像素误差。已验证的像素目标可保持8秒用于云台稳定，但报告仍必须满足当前照片和运动证据的新鲜度。YOPO在这里提供运动基元与轨迹代价的设计启发，实际代价由代码枚举计算，未由神经网络预测。[2][3]')

page('5 严格短轨迹坐标上报')
p('v22取得分数的关键改动是将坐标报告与协同导航轨迹确认解耦。视觉已经给出可用坐标时，不必继续等待导航轨迹累计4个地理身份样本；但报告必须通过更完整的像素连续性、类别、运动和时空检查。旧的地理轨迹报告分支已关闭，地理轨迹仍用于飞行协同。')
table(['检查项','v22实际条件'],[
('连续跟踪','至少4帧像素关联；最近至少3帧保持高置信真车身份'),
('外观身份','当前及连续身份帧的真车概率≥0.95，超过另两类的分差≥0.9'),
('连续定位','至少2次相邻有效估计，间隔不超过1秒；相邻水平位移≤10+35Δt米'),
('图像运动','至少2次连续相对背景运动证据；平面换算速度3.5至18米/秒'),
('定位姿态','工程预算≤110米；生成快速报告估计时俯角至少70度'),
('时空绑定','报告估计必须是当前有效估计；照片哈希一致，接收年龄在0至0.6秒内'),
('执行条件','本机UID一致且状态为active；报告开关开启；共用1Hz限频')],[3.1,13.3])
h('失败原因与修复依据',2)
p('旧版本存在两类失败。第一类是把树木或其他背景识别成真车后上报，v15和v19各有一次大偏差报告，累计RMSE分别约1577米和1372米。单纯降低置信度无法解决这类问题，必须同时检查运动与定位一致性。')
p('第二类是连续真车观测被等待窗口阻断。v18中曾出现连续四帧高置信识别和较小预算定位，但云台随后从0度跳到约47度，定位中断，导航轨迹未准备好。v19改善关联和步幅后，又出现首帧置信度较弱、后三帧很强的情况；等足四帧高置信时，云台运动已使定位预算升高。')
p('v22因此保留至少四帧连续像素轨迹，要求其中最近三帧为高置信身份；允许从两帧高置信且有运动证据的阶段累计几何一致性，并将上报预算设为110米。它没有按目标编号或已知路线决定是否上报。修复的目的，是在证据完整的短窗口内及时发送坐标。')
h('上报与捕获分别验收',2)
p('坐标精度报告与K=2捕获是两个得分通道。报告坐标并不意味着第二架无人机已经看见同一目标；只有引擎判定两机持续有效跟踪同一真实目标20秒，才会清除目标。短中断容忍2秒。当前v22只能据官方结果确认精度得分，不能据报告次数或OBSERVE状态声称协同捕获成功。')

page('6 正式实验与证据边界')
p('实验使用未修改的正式场景，seed为101，请求时长600秒，官方最后记录为599.917秒。冻结提交文件与运行记录的代码及权重哈希一致。最终总分9.06，精度维度为30.2，按该维度0.3权重贡献约9.06分；清除与时间维度均为0。[4]')
table(['版本','seed','报告数','总分','清除数'],[
('v13','101','0','0','0'),('v14','102','0','0','0'),('v15','103','1','0','0'),
('v18','101','0','0','0'),('v19','101','1','0','0'),('v21','101','0','0','0'),('v22','101','2','9.06','0')],[3.1,2.8,3.5,3.5,3.5])
picture('official-score.png')
p('图2 原始官方评分时间序列  运行中9.47分更新为最终9.06分','Caption')
p('两次报告均由严格短轨迹分支发出，报告定位RMSE为11.269米。三机推理异常和异步丢弃均为0，本回合惩罚为0；官方三个目标的coop_ticks均为0，未形成同目标双机有效跟踪。软件成功执行与任务完整完成必须分别评价。')
p('隔离检查验证了三份私有模型及禁止回调文件I/O条件下的推理；受控真车照片回放产生12次报告，诱饵回放为0。回放使用合成队友广播，动作不影响后续照片，因此只是链路检查。10项报告与关联检查、17项视觉检查、31项基础检查也不能替代正式裁判。60个受检官方文件哈希未改变。[6]')
p('表中的版本同时改变过感知、控制或门限，且部分seed不同；即使seed相同，实际采样和闭环轨迹也会变化。这些是工程迭代记录，不是严格控制变量的消融实验。当前只能确认一次非零正式结果，尚未建立多seed统计、未见场景泛化或稳定高分结论。')

page('7 部署与复现')
p('工程根目录为D:\\catkin_ws\\hf2026-sim-windows\\ZqhjGame。正式提交目录为artifacts/submission/score-v22，其中agent.py与vision.pt须保持在一起；requirements.txt记录视觉依赖，baseline_evaluation.json保存原始结果副本。以下命令在发行包父目录的PowerShell中执行，每次使用新的输出目录。')
h('重新训练外观模型',2)
code('Set-Location D:\\catkin_ws\\hf2026-sim-windows\n.\\ZqhjGame\\.venv-learning\\Scripts\\python.exe -B `\n  ZqhjGame/tools/train_vehicle_appearance.py `\n  --reviewed ZqhjGame/artifacts/vision/datasets/appearance-data-v3/accepted.jsonl `\n  --output ZqhjGame/artifacts/vision/models/appearance-rebuild `\n  --steps 2400')
p('该命令读取已有审核图像及标签，不启动比赛。默认不启用后续颜色增强实验。重新训练的权重必须重新验证，不能继承当前冻结包的9.06分。')
h('导出独立Agent',2)
code('.\\ZqhjGame\\vision.cmd export `\n  --weights ZqhjGame/artifacts/vision/models/appearance-rebuild/appearance.pt `\n  --patch-appearance --geometry estimated --distributed-search `\n  --reports-default-on `\n  --output ZqhjGame/artifacts/submission/first-score-rebuild')
p('vision.pt是局部分类网络state_dict，因此需要--patch-appearance。--reports-default-on让SDK直接加载EntryAgent时也开启有门限的报告。')
h('运行已经实测的v22包',2)
code('.\\ZqhjGame\\vision.cmd run --ue-direct --duration 600 --seed 104 `\n  --submission ZqhjGame/artifacts/submission/score-v22/agent.py `\n  --enable-reports --max-photos 500 `\n  --output ZqhjGame/artifacts/vision/runs/score-v22-retest104')
p('该命令直接使用已经验证的权重，无需重新训练。seed104是后续复测示例，本报告未宣称已取得该seed的成绩。不要同时启动两场UE/Redis比赛。')
h('回合结束后读取结果',2)
code('.\\ZqhjGame\\.venv-learning\\Scripts\\python.exe -B `\n  ZqhjGame/tools/analyze_score_run.py `\n  ZqhjGame/artifacts/vision/runs/score-v22-retest104 `\n  --output ZqhjGame/artifacts/checks/score-v22-retest104.json')
p('原始评分文件位于运行目录official/*.evaluation.json。完整训练、隔离检查和推理命令另见项目FIRST_SCORE.md；新机器环境配置见VISION_TRAINING_AND_INFERENCE.md。')

page('8 局限 后续方向与依据')
h('当前适用边界',2)
p('当前方法优先保证少量报告有较充分的视觉依据。高置信身份和速度门限可能漏掉慢车、遮挡车辆或低对比度目标；全局单应变换也可能把不同高度的背景视差当成物体运动。局部平面和接收位姿尚未完成全场景标定，报告误差仍可能随地形和云台变化而增大。')
p('当前安全规划只使用公开边界、已知威胁区和合法队友状态，没有构建完整的视觉深度地图或ESDF，也没有复现论文中的多障碍环境感知。零惩罚是这一次回合的观测结果。协同分工仍采用简单的稳定排序与超时策略，未实现基于共同可见性、到达时间和遮挡预测的队友选择。')
h('后续优化顺序',2)
p('后续应先保留v22作为可复现基线，补齐未见布景的真车与诱饵验证，分开统计漏检、误识和坐标误差；再优化相机时间对齐与局部地形估计。协同阶段应围绕同目标双机实际有效跟踪时长设计会合、视线保持和退出策略，以官方捕获计时验收。')
p('如果恢复YOPO式学习主线，需要把视觉或时序编码器、固定翼候选轨迹及任务代价接入统一训练图，验证任务梯度到达视觉参数，并处理双机可见性与真假识别。仅把独立分类器接到轨迹网络前面，或提高mAP，不能证明已经实现端到端引导学习。此项为后续方向，不属于当前9.06分版本的已完成功能。')
h('参考文献与工程证据',2)
p('[1] Gou F, Du H, Zhao C, Cai Y. A Policy-Guided Reinforcement Learning Method for Encirclement Control in Multiobstacle Environment. IEEE Transactions on Neural Networks and Learning Systems, 36(9), 2025. DOI 10.1109/TNNLS.2025.3566548。')
p('[2] Lu J, Zhang X, Shen H, Xu L, Tian B. You Only Plan Once: A Learning-Based One-Stage Planner With Guidance Learning. IEEE Robotics and Automation Letters, 9(7), 2024. DOI 10.1109/LRA.2024.3399589。')
p('[3] 冻结实现：artifacts/submission/score-v22/agent.py。主要模块包括zqhj_patch_vision、zqhj_visual_geometry、zqhj_photo_entry、zqhj_score_search、zqhj_planner与zqhj_cooperation。')
p('[4] 官方原始评分：artifacts/vision/runs/score600-v22-seed101/official/coop_decoy_1789108334.evaluation.json。提交包中的baseline_evaluation.json与其哈希相同。')
p('[5] 训练记录：artifacts/vision/models/appearance-v3/training.json；审核数据：artifacts/vision/datasets/appearance-data-v3/accepted.jsonl。')
p('[6] 验证记录：artifacts/checks/score-v22-isolated.json、score-v22-101-analysis.json、score-v22-release-integrity.json；规则适配依据见docs/SDK_CONTRACT.md及官方实现。以上工程路径均相对ZqhjGame。')
p('代码 SHA256');code(M['agent_sha256'])
p('权重 SHA256');code(M['vision_sha256'])

path=OUT/'ZqhjGame_v22_技术报告.docx';doc.save(path)
print(path)
